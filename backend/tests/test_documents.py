import unittest
import os
import io
import tempfile
import fitz
from docx import Document
from unittest.mock import patch, MagicMock
from fastapi import HTTPException
from fastapi.security import HTTPAuthorizationCredentials
from fastapi.testclient import TestClient

from app.main import app
from app.auth.firebase_auth import get_current_user
from app.utils.file_validator import validate_file_security, temp_document_file, FileValidationError
from app.services.document_parser import DocumentParserService, TextExtractionError

class TestDocumentSecurityAndExtraction(unittest.TestCase):

    def setUp(self):
        self.client = TestClient(app)

    # 1. Unauthenticated Upload Test
    def test_unauthenticated_upload(self):
        """Test uploading a file without Firebase auth token returns 401."""
        response = self.client.post(
            "/api/documents/upload",
            files={"file": ("contract.pdf", b"%PDF-1.4 test content", "application/pdf")}
        )
        self.assertEqual(response.status_code, 401)

    # 2. Valid PDF Upload Validation & Extraction
    def test_valid_pdf_extraction(self):
        """Test valid PDF parsing extracts text and page numbers."""
        doc = fitz.open()
        page = doc.new_page()
        page.insert_text((50, 50), "SECTION 1. CONFIDENTIALITY CLAUSE\nBoth parties agree to protect secret information.")
        pdf_bytes = doc.tobytes()
        doc.close()

        with temp_document_file(pdf_bytes, ".pdf") as temp_path:
            result = DocumentParserService.extract_from_pdf(temp_path)
            self.assertEqual(result["file_type"], "pdf")
            self.assertEqual(result["page_count"], 1)
            self.assertIn("CONFIDENTIALITY CLAUSE", result["full_text"])

    # 3. Valid DOCX Upload Validation & Extraction
    def test_valid_docx_extraction(self):
        """Test valid DOCX parsing extracts text and section styles."""
        doc = Document()
        doc.add_heading("EMPLOYMENT AGREEMENT", level=1)
        doc.add_paragraph("Employee agrees to a 12-month non-compete clause.")
        
        buffer = io.BytesIO()
        doc.save(buffer)
        docx_bytes = buffer.getvalue()

        with temp_document_file(docx_bytes, ".docx") as temp_path:
            result = DocumentParserService.extract_from_docx(temp_path)
            self.assertEqual(result["file_type"], "docx")
            self.assertIn("non-compete", result["full_text"])

    # 4. Invalid Extension Check
    def test_invalid_extension(self):
        """Test file with unsupported extension (e.g. .exe or .txt) is rejected."""
        mock_file = MagicMock()
        mock_file.filename = "malware.exe"
        mock_file.content_type = "application/x-msdownload"
        
        with self.assertRaises(FileValidationError) as cm:
            validate_file_security(mock_file, b"executable bytes")
        self.assertIn("Unsupported file extension", cm.exception.detail)

    # 5. Invalid MIME Type Check
    def test_invalid_mime_type(self):
        """Test valid extension with forbidden MIME type is rejected."""
        mock_file = MagicMock()
        mock_file.filename = "contract.pdf"
        mock_file.content_type = "image/jpeg"
        
        with self.assertRaises(FileValidationError) as cm:
            validate_file_security(mock_file, b"%PDF-1.4 test")
        self.assertIn("Invalid MIME type", cm.exception.detail)

    # 6. Invalid Magic Bytes Check
    def test_invalid_magic_bytes(self):
        """Test fake PDF file (renamed .txt or malware) failing magic bytes signature is rejected."""
        mock_file = MagicMock()
        mock_file.filename = "fake_contract.pdf"
        mock_file.content_type = "application/pdf"
        
        with self.assertRaises(FileValidationError) as cm:
            validate_file_security(mock_file, b"THIS IS NOT A REAL PDF FILE")
        self.assertIn("File header does not match PDF signature", cm.exception.detail)

    # 7. Oversized File Check (Exceeding 4 MB)
    def test_oversized_file(self):
        """Test file > 4 MB is rejected with size limit error."""
        mock_file = MagicMock()
        mock_file.filename = "huge_contract.pdf"
        mock_file.content_type = "application/pdf"
        
        large_bytes = b"%PDF" + b"0" * (4 * 1024 * 1024 + 100)
        with self.assertRaises(FileValidationError) as cm:
            validate_file_security(mock_file, large_bytes)
        self.assertIn("exceeds maximum limit of 4 MB", cm.exception.detail)

    def test_valid_size_file(self):
        """Test file <= 4 MB passes size validation."""
        mock_file = MagicMock()
        mock_file.filename = "valid_contract.pdf"
        mock_file.content_type = "application/pdf"
        
        valid_bytes = b"%PDF" + b"0" * (3 * 1024 * 1024)
        # Should not raise exception
        validate_file_security(mock_file, valid_bytes)

    # 8. Malformed PDF Handling
    def test_malformed_pdf(self):
        """Test corrupted PDF file passes magic bytes but fails parsing cleanly."""
        corrupted_pdf = b"%PDF-1.4 CORRUPTED DATA TRAILER REMOVED"
        with temp_document_file(corrupted_pdf, ".pdf") as temp_path:
            with self.assertRaises(TextExtractionError) as cm:
                DocumentParserService.extract_from_pdf(temp_path)
            self.assertIn("Failed to parse PDF file", cm.exception.detail)

    # 9. Malformed DOCX Handling
    def test_malformed_docx(self):
        """Test corrupted DOCX file passes magic bytes but fails python-docx parsing cleanly."""
        corrupted_docx = b"PK\x03\x04 CORRUPTED ZIP STREAM"
        with temp_document_file(corrupted_docx, ".docx") as temp_path:
            with self.assertRaises(TextExtractionError) as cm:
                DocumentParserService.extract_from_docx(temp_path)
            self.assertIn("Failed to parse DOCX file", cm.exception.detail)

    # 10. Temporary File Cleanup Guarantee
    def test_temporary_file_cleanup(self):
        """Test that temporary local file is guaranteed to be deleted after context exit."""
        temp_path_stored = None
        with temp_document_file(b"%PDF-1.4 test content", ".pdf") as temp_path:
            temp_path_stored = temp_path
            self.assertTrue(os.path.exists(temp_path))

        # After exiting context manager, temp file MUST be deleted
        self.assertFalse(os.path.exists(temp_path_stored))

    # 11. User Ownership & Token Extraction Integration
    @patch("app.api.documents.FirestoreService.save_document_meta")
    @patch("app.auth.firebase_auth.auth.verify_id_token")
    def test_authenticated_upload_endpoint(self, mock_verify, mock_save_meta):
        """Test end-to-end endpoint with valid token extracts verified UID and writes Firestore metadata."""
        mock_verify.return_value = {
            "uid": "verified_user_uid_123",
            "email": "sarah@legallens.ai"
        }

        # Create minimal valid PDF
        doc = fitz.open()
        page = doc.new_page()
        page.insert_text((50, 50), "Test Agreement Document")
        pdf_bytes = doc.tobytes()
        doc.close()

        response = self.client.post(
            "/api/documents/upload",
            headers={"Authorization": "Bearer mock_valid_token"},
            files={"file": ("sample_agreement.pdf", pdf_bytes, "application/pdf")}
        )

        self.assertEqual(response.status_code, 200)
        res_data = response.json()
        self.assertEqual(res_data["status"], "success")
        self.assertEqual(res_data["meta"]["uid"], "verified_user_uid_123")
        self.assertEqual(res_data["meta"]["originalFileName"], "sample_agreement.pdf")
        self.assertEqual(res_data["meta"]["fileType"], "pdf")
        
        # Verify metadata saved under user's verified UID
        mock_save_meta.assert_called_once()
        saved_uid = mock_save_meta.call_args.kwargs["uid"]
        self.assertEqual(saved_uid, "verified_user_uid_123")

if __name__ == "__main__":
    unittest.main()
