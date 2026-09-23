import fitz  # PyMuPDF
from docx import Document
from typing import List, Dict, Any
from fastapi import HTTPException

class TextExtractionError(HTTPException):
    def __init__(self, detail: str):
        super().__init__(status_code=422, detail=detail)

class DocumentParserService:
    """Service to parse text, page numbers, and structural sections from PDF and DOCX files."""

    @staticmethod
    def extract_from_pdf(file_path: str) -> Dict[str, Any]:
        """Extract text page-by-page from PDF file using PyMuPDF."""
        try:
            doc = fitz.open(file_path)
            pages: List[Dict[str, Any]] = []
            full_text_chunks: List[str] = []
            total_pages = len(doc)

            for page_index in range(total_pages):
                page = doc.load_page(page_index)
                text = page.get_text("text").strip()
                if text:
                    pages.append({
                        "page": page_index + 1,
                        "text": text
                    })
                    full_text_chunks.append(f"--- PAGE {page_index + 1} ---\n{text}")

            doc.close()

            if not pages:
                raise TextExtractionError("PDF file contains no extractable text (it may be scanned images without OCR).")

            return {
                "file_type": "pdf",
                "page_count": total_pages,
                "extracted_pages": pages,
                "full_text": "\n\n".join(full_text_chunks),
                "total_characters": sum(len(p["text"]) for p in pages)
            }
        except fitz.FileDataError:
            raise TextExtractionError("Failed to parse PDF file. The file appears corrupted or invalid.")
        except Exception as e:
            if isinstance(e, HTTPException):
                raise e
            raise TextExtractionError(f"Error during PDF text extraction: {str(e)}")

    @staticmethod
    def extract_from_docx(file_path: str) -> Dict[str, Any]:
        """Extract text, headings, and paragraph sections from DOCX file using python-docx."""
        try:
            doc = Document(file_path)
            sections: List[Dict[str, Any]] = []
            full_text_chunks: List[str] = []
            current_section_title = "Preamble"

            for para in doc.paragraphs:
                text = para.text.strip()
                if not text:
                    continue
                
                # Check for headings
                style_name = para.style.name.lower() if para.style else ""
                if "heading" in style_name or para.style.name.startswith("Heading"):
                    current_section_title = text
                
                sections.append({
                    "section": current_section_title,
                    "style": para.style.name if para.style else "Normal",
                    "text": text
                })
                full_text_chunks.append(text)

            # Extract table text if present
            table_count = len(doc.tables)
            for table in doc.tables:
                for row in table.rows:
                    row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                    if row_text:
                        full_text_chunks.append(row_text)

            if not full_text_chunks:
                raise TextExtractionError("DOCX file contains no extractable text.")

            return {
                "file_type": "docx",
                "page_count": max(1, len(full_text_chunks) // 30), # Estimated equivalent pages
                "table_count": table_count,
                "extracted_sections": sections,
                "full_text": "\n\n".join(full_text_chunks),
                "total_characters": sum(len(t) for t in full_text_chunks)
            }
        except Exception as e:
            if isinstance(e, HTTPException):
                raise e
            raise TextExtractionError(f"Failed to parse DOCX file: {str(e)}")
